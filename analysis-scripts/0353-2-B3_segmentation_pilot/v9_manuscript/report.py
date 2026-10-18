"""Editable Plant Phenomics-oriented manuscript materials, without invented rules."""
import argparse,re
from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import common as c

def mdtable(headers,rows):return '|'+ '|'.join(headers)+'|\n|'+'|'.join(['---']*len(headers))+'|\n'+'\n'.join('|'+ '|'.join(map(str,r))+'|' for r in rows)
def wordtable(doc,headers,rows):
 t=doc.add_table(rows=1,cols=len(headers));t.style='Light Shading Accent 1'
 for cell,v in zip(t.rows[0].cells,headers):cell.text=str(v)
 for row in rows:
  for cell,v in zip(t.add_row().cells,row):cell.text=str(v)
 return t
def newdoc(title):
 d=Document();sec=d.sections[0];sec.top_margin=sec.bottom_margin=Inches(.7);style=d.styles['Normal'];style.font.name='Arial';style.font.size=Pt(10);style._element.rPr.rFonts.set(qn('w:eastAsia'),'宋体');d.add_heading(title,0);return d

def main():
 p=argparse.ArgumentParser();p.add_argument('--run',type=Path,required=True);a=p.parse_args();run=a.run;out=run/'manuscript';out.mkdir(exist_ok=True);j=c.read(run/'tables/summary.json');agg=j['aggregate'];lens=j['length_agreement'];caps=c.read(run/'figures/captions.json');design=c.read(run/'tables/study_design.json')
 def row(n,method):return next(r for r in agg if r['model']==n and r['aggregation']==method)
 primary=[[c.NAMES[n]]+[f'{row(n,"C4_pooled4dates")[k]:.4f}' for k in c.METRICS] for n in c.MODELS]
 validation=[[c.NAMES[n]]+[f'{row(n,"C3_pooled4dates")[k]:.4f}' for k in c.METRICS] for n in c.MODELS]
 meanrows=[[c.NAMES[n],f'{row(n,"C4_period_macro_mean")["Dice"]:.4f} ± {row(n,"C4_period_macro_SD")["Dice"]:.4f}',f'{row(n,"C3_C4_arithmetic_region_mean")["Dice"]:.4f}',f'{row(n,"C3_C4_count_pooled8views")["Dice"]:.4f}'] for n in c.MODELS]
 lrows=[]
 for n in c.MODELS:
  m=next(r for r in lens if r['model']==n and r['region']=='C4');lrows.append([c.NAMES[n],f'{m["R2_agreement"]:.4f}',f'{m["RMSE_px"]:.2f}',f'{m["MAE_px"]:.2f}',f'{m["CCC"]:.4f}',f'{m["bias_px"]:.2f}'])
 headers=['Model','Precision','Recall','Dice','IoU'];lheaders=['Model','R² agreement','RMSE (px)','MAE (px)','CCC','Bias (px)']
 methods=('A pilot root/background segmentation analysis used the four upper-period RGB images from minirhizotron position0353-2-B3 (2541 ×1108 native pixels each). The four images are repeated observations of one location, rather than independent biological replicates. The original human progress file and binary masks were retained without alteration. Uncertain pixels were excluded as originally encoded; unpainted upper-image background followed the earlier annotation-completion declaration. No temporal-state or death classes were trained in this v9 analysis. '
 'Five448-pixel-wide spatial columns (C0–C4) began at x=22,534,1046,1558 and2070, with64-pixel guards. C0–C2 were used for training, C3 for checkpoint/threshold/postprocessing selection, and C4 for a repeatedly inspected same-position benchmark. All dates from a spatial column were assigned to the same group. These correspond to12 training,4 validation and4 benchmark date-region views derived from4 original images; no independent external test images were available. Each model used256 fine-tuning crop draws (189 distinct date/coordinate windows), with40 validation and40 benchmark evaluation crops per model. Native inference context differed (384 for U-Net,256 for the other models), and original pretraining and optimizers also differed; this is not an equal-total-compute architecture ranking. Formal scores used fused native columns with every valid pixel counted once. '
 'U-Net, the uploaded TensorFlow1.12 DeepLabV3+ and SegFormer-B0 were compared using the frozen v9 networks and postprocessing. Four-period pooled Precision, Recall, Dice and IoU were calculated after summing TP,FP andFN across dates within each evaluation region. Equal-weighted date means and sample SD were reported separately as descriptive temporal variation, not biological uncertainty. C3 and C4 were not pooled for the headline score. Root-length agreement used skeleton-derived native-pixel lengths in eight224×277-pixel blocks per region across four dates. R² was defined as1−SSE/SST relative to the identity relation; RMSE,MAE,CCC and descriptive Bland–Altman bias±1.96SD were also reported. These32 block-date pairs per region are correlated. No p-values or independent-biological confidence intervals were inferred. Raw-score average precision grouped equal-score ties and used all valid pixels. This re-analysis did not retrain the models or change the masks, thresholds or annotations.')
 results=('On the repeatedly inspected same-position C4 benchmark, U-Net achieved foreground Precision=0.6197, Recall=0.6502, Dice=0.6346 andIoU=0.4648. DeepLabV3+ achieved0.5100,0.5317,0.5206 and0.3519, respectively, whereas SegFormer-B0 achieved0.6342,0.4431,0.5217 and0.3529. Thus,U-Net showed the highest Dice andIoU in this pilot comparison; SegFormer showed slightly higher precision but substantially lower recall. The C3 validation Dice values were0.7538,0.6573 and0.7087 for U-Net,DeepLabV3+ andSegFormer, respectively. The corresponding C3-to-C4 decreases were0.1193,0.1366 and0.1870, indicating sensitivity to the evaluated spatial region rather than uniformly high segmentation accuracy. These differences do not establish a causal effect of a specific background artifact. '
 'For U-Net,the equal-weighted C4 date mean Dice was0.6310±0.0716 (descriptive SD), distinct from the pooled-count Dice of0.6346. Combining the C3 andC4 Dice by arithmetic averaging would give0.6942,whereas summing their pixel counts would give0.7176. Both mixed summaries include selection validation and are excluded from the main evaluation claim. '
 'Against human-mask-derived skeleton lengths on C4,U-Net achieved R² agreement=0.8843,RMSE=139.58px andCCC=0.9365. DeepLabV3+ achieved R²=0.6607 andRMSE=239.04px,andSegFormer achieved R²=0.4992 andRMSE=290.42px. Mean length biases were−63.42,−124.36 and−186.96px, respectively. Relatively high length agreement alongside moderate pixel Dice is possible because aggregate skeleton length does not fully measure pixel alignment,root width or root-instance identity. These pixel-unit reference lengths were derived from masks and were not independent manual length measurements. '
 'The results support a limited feasibility assessment at the annotated position. They do not demonstrate Dice above0.9,significant differences among architectures,or generalization to other positions or tubes. Four-period whole-field maps include training regions and are qualitative illustrations,not independent accuracy estimates.')
 limitations=('This pilot uses one annotated location. C3 served model selection,andC4 has been repeatedly inspected during development. Consequently,C4 should be described as a same-position benchmark rather than a blind external test. Spatial guards reduce direct crop overlap but do not create independent plants or tubes. Pretraining,inference context and optimization histories differ across models. Human labels are not an independently adjudicated biological gold standard,androot visibility,thin-root ambiguity and mask width can affect pixel scores. A small separately annotated set from previously unused positions or tubes is needed before claiming transfer accuracy; annotation of every production image is not required. Binary foreground predictions cannot establish root birth,extension,persistence,browning or death. No image scale was available for conversion from pixels to physical root length.')
 availability=('Original RGB,annotations,trained v9 result bundles and reproducible source code are available from Annalee811/LI at commit608659941cd720a8292efbebda782e5277e30f08. This publication re-analysis is saved under v9-manuscript-ready-v1 with input/archive SHA256 manifests. The original annotation SHA256 is27943e5f9d087bf718668bba0a39058dc9a4871de6c75c5ccc27febad4c06a09. If a persistent repository DOI is issued,insert the verified DOI and version here; do not invent a DOI or describe a mutable branch as immutable. No journal submission orDOI registration is performed by this report.')
 def english(text):
  text=re.sub(r'([a-z])(?=\d)',r'\1 ',text);text=re.sub(r'(\d)(?=[a-z])',r'\1 ',text);text=re.sub(r',(?=[A-Za-z])',', ',text);text=re.sub(r'\b(and|or|for)(?=[A-Z])',r'\1 ',text);return text
 methods,results,limitations,availability=map(english,[methods,results,limitations,availability])
 methods=methods.replace('All dates from a spatial column were assigned to the same group.','All dates from a spatial column were assigned to the same group (Fig. 1).')
 results=results.replace('The C3 validation Dice values were','The primary C4 comparison is shown in Fig. 2 and Table 1. Representative fixed-field errors across all four dates are visualized in Fig. 3. The C3 validation Dice values were').replace('Both mixed summaries include selection validation','These alternative aggregations are shown in Fig. S4, and date means with descriptive SD in Fig. S6. Both mixed summaries include selection validation').replace('Relatively high length agreement alongside','The length comparisons are shown in Fig. 4 and Table 2. Relatively high length agreement alongside')
 draft=newdoc('Pilot root segmentation in longitudinal minirhizotron images — editable results draft')
 draft.add_paragraph('Prepared for Plant Phenomics consideration. This is a methods/results draft,not a complete submitted manuscript. The journal author-guideline website could not be accessed from this environment on18October2026; exact current submission requirements remain unverified.')
 for title,text in [('Methods: pilot design and evaluation',methods),('Results: segmentation and root-length agreement',results),('Limitations and appropriate claims',limitations),('Data and code availability — verify archive version before submission',availability)]:draft.add_heading(title,1);draft.add_paragraph(text)
 draft.add_heading('Table1. Primary pilot evaluation: C4 pooled over four dates',1);wordtable(draft,headers,primary);draft.add_paragraph('C4 is a repeated same-position benchmark,not an independent external test. Metrics are strict native-pixel foreground ratios.')
 draft.add_heading('Table2. C4 root-length agreement',1);wordtable(draft,lheaders,lrows);draft.add_paragraph('32 correlated block-date pairs per model;human-mask-derived reference lengths,native pixels;no independent manual measurement or significance claim.')
 draft.add_heading('TableS1. Selection validation C3',1);wordtable(draft,headers,validation)
 draft.add_heading('Figure legends',1)
 for name,cap in sorted(caps.items()):draft.add_heading(name,2);draft.add_paragraph(cap)
 draft.add_heading('Information the authors must add from actual records',1)
 draft.add_paragraph('Add the actual image-acquisition site and dates,plant material and growth conditions,minirhizotron tube/camera specifications,physical calibration if available,annotation operator and procedure,andappropriate model/method references. Author names,affiliations,funding,contributions and declarations must be entered from actual records. Do not infer these from image filenames.')
 draft.save(out/'Manuscript_Methods_Results_EN.docx');(out/'Manuscript_Methods_Results_EN.md').write_text('# Methods\n\n'+methods+'\n\n# Results\n\n'+results+'\n\n# Limitations\n\n'+limitations+'\n\n# Data and code availability\n\n'+availability+'\n')
 plan=[('Fig1','正文Methods或Results开头','空间划分与四期同组；4张原图不写成20张独立原图'),('Fig2','正文主要性能图','C4四期像素计数合并指标；小空心点展示逐期表现'),('Fig3','正文定性验证','固定C4区域，同一位置比较三模型漏检/误检'),('Fig4','正文表型一致性','三模型根长回归+Bland–Altman；单位像素'),('Table1','正文主要表','C4 Precision、Recall、Dice、IoU'),('Table2','正文或补充','C4根长R²、RMSE、MAE、CCC、偏差'),('FigS1 / TableS1','补充材料，或与正文C4图并列但分别标注','C3用于选优，不能代替C4'),('FigS2','补充材料','C3根长一致性，与C4分开展示'),('FigS3','补充材料','四期、两区差异与人工根像素占比'),('FigS4','内部汇总或补充，非核心成绩','应用户要求给出不同平均/合并口径；不能冒充独立测试'),('FigS5','补充材料','PR曲线与准确AP，前景稀疏时比总准确率更有解释性'),('FigS6','补充材料','C3、C4各自四期均值±SD，仅描述时间波动')]
 cn='''# Plant Phenomics：v9结果重整与使用说明

本次只重新核算和排版用户指定的v9根/背景二分类成果，不训练模型、不改人工标注、不改变预测掩膜；不混入v13时序标签修订。已逐字节核对57个原始RGB/掩膜/概率/标注文件与原v9归档一致。

## C3、C4是什么？

C0–C4是同一位置上半部RGB沿水平方向划出的空间列，不是五个模型、五根管或五个时期。C0–C2训练；C3验证，选择模型轮次、阈值及连接后处理；C4为反复查看过的同位置评价区。原图2541×1108，列宽448，相邻留64像素间隔；C3为x1558–2005、C4为x2070–2517（0起始像素坐标）。同一列的T1–T4始终同组。

这种分组能避免同一位置前后期随机落入训练和评价，但不能构成跨位置或跨根管独立验证。4张原图得到12/4/4个时期×区域视图；独立外部测试图像为0。256次训练抽样只有189个不同的时期/坐标窗口，不能作为256张独立原图。

## 论文该选哪个值？

推荐正文以C4四期合并指标为本预实验的主要评价，明确称同位置benchmark。C3验证单列补充；若正文篇幅允许，可以C3/C4并列，但不能省略用途。不要只挑较高的C3，也不要把C3+C4平均作为唯一成绩。C4已被反复查看，不能写成盲测、独立测试或外部验证。

论文整体结论仍应定位为单位置预实验。我的投稿判断是：若用于Plant Phenomics稿件，应围绕实际表型问题和可验证结果，而不能仅凭这组分数声称通用高精度系统或跨管泛化。这是研究设计判断，并非声称该刊规定了某个最小样本数。

## 三种平均不能混用

1. 推荐的C4四期合并：先求四期TP、FP、FN之和，再算Dice=2ΣTP/(2ΣTP+ΣFP+ΣFN)。这是主表口径，重叠推理图块每个原像素只计一次。
2. C4四期宏均值：每期先算指标，再对4个值等权平均；可写均值±样本SD，SD描述4个相关时期的变化，不是独立植株的误差或置信区间。
3. C3+C4简单均值、C3+C4像素合并、8个日期区域等权均值均为混合验证区与评价区的探索性汇总，已经作图，但不作为主结论。R²不能用两个分区R²简单平均；合并描述必须把原配对数据拼接后重新计算。

## 正文主要分割结果（严格原像素指标）

'''+mdtable(headers,primary)+'\n\n## C3验证结果（单独报告）\n\n'+mdtable(headers,validation)+'\n\n## 已计算的均值及混合汇总\n\n'+mdtable(['模型','C4四期Dice均值±SD','C3/C4 Dice简单均值','C3+C4像素合并Dice'],meanrows)+'\n\n混合像素合并会给根像素及错误计数较多的区域更大权重。例如U-Net从C4的0.6346变成混合汇总的0.7176，不代表模型改进，而是换了被评价的数据范围和权重。\n\n## 根长一致性，单位仅为原像素\n\n'+mdtable(lheaders,lrows)+'''

较高R²不能代替较高Dice：骨架总长度可能接近，但根的位置、宽度或局部分割仍有偏差。这里R²=1−SSE/SST相对于y=x，不是OLS线的R²或相关系数平方；图中实线OLS、虚线恒等线。参考来自人工掩膜骨架，不是另一套独立人工测长。每区8块×4期=32对相关观测，不当作32个生物重复。Bland–Altman的±1.96SD是描述性一致性界限，不是独立生物置信区间。

## 推荐放哪些图表？

'''+mdtable(['材料','建议位置','用途'],plan)+'''

原v9四期完整场景图含训练区，只能作整体展示，不能声称是四张独立测试图。固定ROI图不替代整列评价。没有对某些根不清楚的时期或区域进行删选。

## 中文结果分析示例（按实际指标，可自行改写）

在0353-2-B3位置的同位置C4评价区，U-Net获得最高的Dice（0.6346）和IoU（0.4648），Precision和Recall分别为0.6197和0.6502。DeepLabV3+与SegFormer-B0的Dice分别为0.5206与0.5217。SegFormer-B0的Precision略高于U-Net，但Recall较低，表明其在本数据中的保守预测伴随着更多漏检。三模型从C3到C4的Dice均下降，提示性能受到空间背景与根系分布差异影响；当前结果不足以将差异归因于特定水滴或土壤类型。

U-Net预测根长与人工掩膜衍生骨架长度在C4上的一致性R²为0.8843，RMSE为139.58像素。DeepLabV3+和SegFormer-B0的R²分别为0.6607、0.4992。较高的长度一致性与中等像素分割指标并存，说明长度指标不能取代对空间掩膜的评价。所有根长结果仅作相关块和时期的描述性分析，不能据此声称显著一致或跨根管泛化。

建议结论为“提供了该位置条件下的初步可行性证据，但精细分割及迁移适用性仍需验证”，不能写“Dice达0.956”“显著优于其他模型”“充分证明可准确预测其他根管”。不要求全部位置都重标；后续只需从未用于开发的位置/根管抽取小规模独立人工核查集，才有依据估计迁移误差。

## Plant Phenomics 投稿文件组织建议

实际上传顺序：先补全英文主稿→从期刊官方主页进入当前投稿入口→按实际研究选择稿件类型→上传Word主稿和独立TIFF主图→按系统要求上传补充材料→填写数据/代码声明并关联归档链接→检查系统生成的预览PDF、图号、字体及作者信息→由作者确认后提交。没有替你登录或提交。

作者指南入口：https://spj.science.org/page/plantphenomics/for-authors
期刊主页：https://spj.science.org/journal/plantphenomics

2026-10-18在当前环境因域名访问限制未能打开这两个官方页面。因此以下为通用准备方案，不冒称已核对该刊最新硬性标准；投稿时请以官网和系统当前要求为准。本次未提交任何稿件。

- 正文：使用Manuscript_Methods_Results_EN.docx中的方法、结果、两张主表和英文图注作为草稿；需补全真实采集方法、研究背景、生物学问题、作者信息和参考文献。不是可直接投稿的完整论文。
- 主图：上传Fig1–Fig4的独立TIFF；同时备份PDF。文件按图号命名、在正文按顺序引用。最终尺寸、字高和文件大小应按该刊要求复核；本包600dpi是排版输出，不代表提高原始光学分辨率。
- 补充材料：FigS1–S6及逐期/汇总CSV放为SupplementaryMaterial，并在正文引用。是否合并为一个PDF、是否需单独上传图注，请按系统字段操作。FigS4不是必放的核心性能图。
- 数据与代码：不能把论文图片ZIP代替数据。保留4张原RGB、原标注、分组/裁剪记录、最终逐像素预测、指标原始计数和脚本。GitHub不可变v9提交已列在数据声明草稿中；若使用Zenodo等归档获得DOI，应在实际取得后补入，不虚构DOI。
- 投稿系统中区分主稿、主图、补充材料、数据/代码链接。系统若要求coverletter，应如实介绍研究问题、增量和局限，不以单位置结果声称完成外部验证。
- 投稿前复核：数据是1位置4时期，不是20张独立原图；C4不是外部盲测；32根长配对不是独立生物样本；二分类结果没有新生/持续/延伸/死亡预测；长度没有物理标尺，不能把px写为mm。

## 完整脚本与原始成果如何恢复？

新分析数据脚本包含可重算所有表图的native输入、SHA256及原始v9完整训练/推理源码快照。原v9七个成果ZIP另按原字节重新保存，保留训练切块、权重、TF运行环境及四期完整输出。所有原档SHA均与原发布版本一致。新分析脚本不依赖GPU，也不重新训练；原网络预测与训练脚本保留自动CUDA及原TF环境说明。
'''
 (out/'论文图表使用与分析_CN.md').write_text(cn)
 d=newdoc('Plant Phenomics：图表使用、C3/C4与投稿说明')
 for block in cn.split('\n\n'):
  if block.startswith('# '):continue
  if block.startswith('## '):d.add_heading(block[3:],1)
  elif block.startswith('|'):
   lines=block.splitlines();table=[[cell.strip() for cell in line.strip('|').split('|')] for line in lines];wordtable(d,table[0],table[2:])
  else:d.add_paragraph(block)
 d.save(out/'Plant_Phenomics_图表与投稿说明_CN.docx')
 d=newdoc('Figure legends — v9 pilot re-analysis')
 for name,cap in sorted(caps.items()):d.add_heading(name,1);d.add_paragraph(cap)
 d.save(out/'Figure_Legends_EN.docx')
 c.csvout(run/'tables/Table_main_C4_length.csv',[r for r in lens if r['region']=='C4'])
 c.js(out/'journal_requirement_status.json',dict(journal='Plant Phenomics',checked_date_client='2026-10-18',official_URLs=['https://spj.science.org/page/plantphenomics/for-authors','https://spj.science.org/journal/plantphenomics'],official_current_guidelines_verified=False,reason='Environment outbound domain restriction: HTTPS proxy CONNECT403',submission_performed=False,file_preparation='General publication defaults:600dpiTIFF,PNG,PDF;quantitativeSVG;editableDOCX;confirm exact current journal requirements before submission'))
 print('EDITABLE_MANUSCRIPT_AND_SUBMISSION_GUIDE_READY',flush=True)
if __name__=='__main__':main()
